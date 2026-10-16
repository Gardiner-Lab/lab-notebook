"""
An export as a Word document (.docx), to edit: made with python-docx.

    Photos      in the text, where they are.
    PDFs        each page as a picture, where the PDF is attached (Word cannot
                hold a PDF's pages as they are), and the PDF itself beside it.
    Other files beside the document: the export is then a zip of the document
                and a files folder, and the document links to each file there.

Headings follow the tree, so Word's navigation pane shows it; the contents at
the start link to each page.
"""

import io
import re
import zipfile
from typing import Callable, Dict, List, Optional, Tuple

from app import exporting, richtext, sheet
from app.exporting import Context, Entry

MAX_PDF_PAGES = 30
MAX_TEXT_LINES = 200
PAGE_WIDTH_IN = 6.3
PAGE_HEIGHT_IN = 8.6
GREY = (0x6B, 0x72, 0x80)
GREEN = (0x04, 0x78, 0x57)


def _docx():
    import docx
    from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Inches, Pt, RGBColor
    return docx, WD_BREAK, WD_COLOR_INDEX, OxmlElement, qn, RT, Inches, Pt, RGBColor


class _Writer:
    def __init__(self, ctx: Context):
        (self.docx, self.WD_BREAK, self.WD_COLOR_INDEX, self.OxmlElement, self.qn, self.RT,
         self.Inches, self.Pt, self.RGBColor) = _docx()
        self.ctx = ctx
        self.doc = self.docx.Document()
        self.files: Dict[str, Optional[dict]] = {}
        self.beside: Dict[str, Tuple[str, bytes]] = {}     # attachment id -> (path in the zip, bytes)
        self.bookmark = 0
        self.mentioned = set()
        style = self.doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = self.Pt(10)
        section = self.doc.sections[0]
        section.left_margin = section.right_margin = self.Inches(0.9)
        section.top_margin = section.bottom_margin = self.Inches(0.8)

    # --- small pieces -------------------------------------------------------------

    def grey(self, paragraph, text: str, size: float = 8.5):
        run = paragraph.add_run(text)
        run.font.size = self.Pt(size)
        run.font.color.rgb = self.RGBColor(*GREY)
        return run

    def link(self, paragraph, text: str, url: Optional[str] = None, anchor: Optional[str] = None):
        """A hyperlink: to an address, or to a bookmark in the document."""
        qn, OxmlElement = self.qn, self.OxmlElement
        link = OxmlElement('w:hyperlink')
        if anchor:
            link.set(qn('w:anchor'), anchor)
        else:
            link.set(qn('r:id'), paragraph.part.relate_to(url, self.RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '1D4ED8')
        underline = OxmlElement('w:u')
        underline.set(qn('w:val'), 'single')
        props.append(color)
        props.append(underline)
        run.append(props)
        t = OxmlElement('w:t')
        t.text = text
        t.set(qn('xml:space'), 'preserve')
        run.append(t)
        link.append(run)
        paragraph._p.append(link)

    def mark(self, paragraph, name: str):
        """A bookmark around the start of a paragraph, for links to point at."""
        qn, OxmlElement = self.qn, self.OxmlElement
        self.bookmark += 1
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(self.bookmark))
        start.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(self.bookmark))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)

    def shade(self, paragraph, fill: str = 'EEF2FF'):
        qn, OxmlElement = self.qn, self.OxmlElement
        props = paragraph._p.get_or_add_pPr()
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:color'), 'auto')
        shading.set(qn('w:fill'), fill)
        props.append(shading)

    def picture(self, paragraph, data: bytes, max_w: float = PAGE_WIDTH_IN, max_h: float = PAGE_HEIGHT_IN):
        """A picture in a paragraph, no wider than the page; formats Word lacks become PNG."""
        from PIL import Image
        try:
            image = Image.open(io.BytesIO(data))
            w, h = image.size
            dpi = image.info.get('dpi', (96, 96))[0] or 96
            if image.format not in ('PNG', 'JPEG', 'GIF', 'BMP'):
                out = io.BytesIO()
                image.convert('RGBA' if image.mode in ('RGBA', 'LA', 'P') else 'RGB').save(out, 'PNG')
                data = out.getvalue()
        except Exception:
            self.grey(paragraph, '[a picture that could not be read]')
            return
        width = min(max_w, w / float(dpi))
        if h and width * h / w > max_h:
            width = max_h * w / h
        paragraph.add_run().add_picture(io.BytesIO(data), width=self.Inches(max(width, 0.3)))

    def _file(self, att_id: str) -> Optional[dict]:
        if att_id not in self.files:
            self.files[att_id] = exporting.load(self.ctx.person, att_id)
        return self.files[att_id]

    def _beside(self, f: dict) -> str:
        """Put a file beside the document, in the zip's files folder; its path there."""
        if f['id'] not in self.beside:
            used = {p for p, _ in self.beside.values()}
            name = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', '_', f['filename']) or 'file'
            path, n = f'files/{name}', 1
            while path in used:
                n += 1
                stem, dot, ext = name.rpartition('.')
                path = f'files/{stem or name} ({n}){dot}{ext if stem else ""}'
            self.beside[f['id']] = (path, f['data'])
        return self.beside[f['id']][0]

    # --- attachments ---------------------------------------------------------------------

    def shown(self, att_id: str) -> None:
        f = self._file(att_id)
        doc = self.doc
        if not f:
            self.grey(doc.add_paragraph(), 'A file that is no longer available.')
            return
        what = f'{exporting.kind_name(f["mime"], f["filename"])}, {exporting.size_label(f["size"])}'
        if f['mime'] in exporting.IMAGE_TYPES:
            self.picture(doc.add_paragraph(), f['data'])
            self.grey(doc.add_paragraph(), f['filename'], 7.5)
            return
        path = self._beside(f)
        p = doc.add_paragraph()
        p.add_run('📎 ')
        self.link(p, f['filename'], url=path)
        if f['mime'] == 'application/pdf':
            reader = exporting.pdf_reader(f['data'])
            pages = len(reader.pages) if reader else 0
            shown = self._pdf_pages(f['data'], min(pages, MAX_PDF_PAGES)) if reader else 0
            note = (f' · {what} · {pages} page{"s" if pages != 1 else ""} below'
                    + (f' (the first {shown})' if 0 < shown < pages else '') if shown else
                    f' · {what} · its pages could not be shown')
            self.grey(p, note + '; the PDF is in the files folder beside this document')
            return
        self.grey(p, f' · {what} · in the files folder beside this document')
        if exporting.is_text(f['mime']):
            lines = exporting.text_of(f['data']).splitlines()
            for line in lines[:MAX_TEXT_LINES]:
                run = doc.add_paragraph().add_run(line)
                run.font.name = 'Consolas'
                run.font.size = self.Pt(8)
            if len(lines) > MAX_TEXT_LINES:
                self.grey(doc.add_paragraph(), f'… {len(lines) - MAX_TEXT_LINES} more lines in the file')

    def _pdf_pages(self, data: bytes, count: int) -> int:
        """Each page of a PDF as a picture in the document; how many were shown."""
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(data)
        except Exception:
            return 0
        shown = 0
        try:
            for k in range(count):
                page = pdf[k]
                image = page.render(scale=1.6).to_pil()
                out = io.BytesIO()
                image.convert('RGB').save(out, 'PNG', optimize=True)
                p = self.doc.add_paragraph()
                self.picture(p, out.getvalue(), max_h=PAGE_HEIGHT_IN - 0.4)
                shown += 1
        except Exception:
            pass
        finally:
            pdf.close()
        return shown

    # --- text -------------------------------------------------------------------------------

    def text(self, block: dict) -> None:
        tree = richtext.parse(block.get('html') or '')
        after: List[str] = []
        self._flow(tree, after, {})
        for att in dict.fromkeys(after):
            self.mentioned.add(att)
            marker = len(self.doc.paragraphs)
            self.shown(att)
            if len(self.doc.paragraphs) > marker:
                self.mark(self.doc.paragraphs[marker], f'a{att[:30]}')

    def _flow(self, node, after: List[str], style: dict, list_stack: Optional[list] = None):
        """Block elements into paragraphs; loose inline content gathers into one."""
        list_stack = list_stack or []
        pending: List = []

        def flush():
            if pending and any(not isinstance(c, str) or c.strip() for c in pending):
                p = self.doc.add_paragraph()
                for c in pending:
                    self._inline(p, c, after, style)
            pending.clear()

        for child in node.children:
            if isinstance(child, richtext.Node) and child.tag in richtext.BLOCKS:
                flush()
                self._block(child, after, style, list_stack)
            else:
                pending.append(child)
        flush()

    def _block(self, node, after, style, list_stack):
        doc = self.doc
        tag = node.tag
        if tag in ('h1', 'h2', 'h3'):
            p = doc.add_paragraph()
            for c in node.children:
                self._inline(p, c, after, {**style, 'b': True, 'size': {'h1': 13, 'h2': 12, 'h3': 11}[tag]})
        elif tag == 'p':
            p = doc.add_paragraph()
            for c in node.children:
                self._inline(p, c, after, style)
        elif tag == 'pre':
            p = doc.add_paragraph()
            run = p.add_run(node.text())
            run.font.name = 'Consolas'
            run.font.size = self.Pt(8.5)
        elif tag == 'blockquote':
            start = len(doc.paragraphs)
            self._flow(node, after, {**style, 'i': True}, list_stack)
            for p in doc.paragraphs[start:]:
                p.paragraph_format.left_indent = self.Inches(0.3)
        elif tag == 'hr':
            self.grey(doc.add_paragraph(), '─' * 40)
        elif tag in ('ul', 'ol'):
            stack = list_stack + [[tag, 0]]
            for c in node.children:
                if isinstance(c, richtext.Node) and c.tag == 'li':
                    self._block(c, after, style, stack)
                elif isinstance(c, richtext.Node):
                    self._block(c, after, style, stack)
        elif tag == 'li':
            level = list_stack[-1] if list_stack else ['ul', 0]
            level[1] += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = self.Inches(0.25 * len(list_stack or [1]))
            p.paragraph_format.first_line_indent = self.Inches(-0.2)
            p.add_run(f'{level[1]}. ' if level[0] == 'ol' else '• ')
            nested = []
            for c in node.children:
                if isinstance(c, richtext.Node) and c.tag in ('ul', 'ol'):
                    nested.append(c)
                elif isinstance(c, richtext.Node) and c.tag in richtext.BLOCKS:
                    for cc in c.children:
                        self._inline(p, cc, after, style)
                else:
                    self._inline(p, c, after, style)
            for n in nested:
                self._block(n, after, style, list_stack)

    def _inline(self, p, node, after, style):
        if isinstance(node, str):
            if not node:
                return
            run = p.add_run(node)
            run.bold = style.get('b') or None
            run.italic = style.get('i') or None
            run.underline = style.get('u') or None
            run.font.strike = style.get('s') or None
            run.font.subscript = style.get('sub') or None
            run.font.superscript = style.get('sup') or None
            if style.get('code'):
                run.font.name = 'Consolas'
            if style.get('size'):
                run.font.size = self.Pt(style['size'])
            if style.get('mark'):
                run.font.highlight_color = self.WD_COLOR_INDEX.YELLOW
            return
        tag, attrs = node.tag, node.attrs
        if tag == 'br':
            p.add_run().add_break()
            return
        if tag == 'img':
            att = attrs.get('data-attachment')
            if att:
                self.mentioned.add(att)
            f = self._file(att) if att else None
            if f and f['mime'] in exporting.IMAGE_TYPES:
                self.picture(p, f['data'])
            else:
                self.grey(p, '[picture not available]')
            return
        if tag == 'a':
            ctx = self.ctx
            text = node.text()
            if attrs.get('data-attachment'):
                f = self._file(attrs['data-attachment'])
                if f:
                    after.append(attrs['data-attachment'])
                    self.link(p, f'📎 {f["filename"]}', anchor=f'a{attrs["data-attachment"][:30]}')
                return
            if attrs.get('data-note'):
                target = attrs['data-note']
                if target in ctx.included:
                    self.link(p, text, anchor=f'n{target[:38]}')
                else:
                    self.link(p, text, url=ctx.note_link(target))
                return
            url = (ctx.plasmid_link(attrs['data-plasmid']) if attrs.get('data-plasmid') else
                   ctx.seed_link(attrs['data-seed']) if attrs.get('data-seed') else
                   ctx.stock_link(attrs['data-stock']) if attrs.get('data-stock') else
                   attrs.get('href') if (attrs.get('href') or '').startswith(('http://', 'https://', 'mailto:')) else None)
            if url and url.startswith(('http://', 'https://', 'mailto:')):
                self.link(p, text, url=url)
            else:
                self._inline(p, text, after, style)
            return
        flags = {'b': 'b', 'strong': 'b', 'i': 'i', 'em': 'i', 'u': 'u', 's': 's', 'sub': 'sub', 'sup': 'sup',
                 'code': 'code', 'mark': 'mark'}
        inner = {**style, flags[tag]: True} if tag in flags else style
        for c in node.children:
            self._inline(p, c, after, inner)

    # --- blocks ----------------------------------------------------------------------------

    def steps(self, block: dict) -> None:
        doc = self.doc
        if block.get('title'):
            doc.add_paragraph().add_run(block['title']).bold = True
        for n, item in enumerate(block.get('items') or [], 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = self.Inches(0.25)
            p.add_run(f'{"☑" if item.get("by") else "☐"} {n}. {item.get("text") or ""}')
            if item.get('by'):
                run = p.add_run(f'   ✓ {item["by"]} · {self.ctx.when(item.get("at"))}')
                run.font.size = self.Pt(8)
                run.font.color.rgb = self.RGBColor(*GREEN)

    def table(self, block: dict) -> None:
        doc = self.doc
        columns = block.get('columns') or []
        rows = block.get('rows') or []
        if block.get('title'):
            doc.add_paragraph().add_run(block['title']).bold = True
        scale = block.get('scale')
        if scale:
            self.grey(doc.add_paragraph(), f'{scale.get("label") or "Reactions"}: '
                      f'{sheet.show(sheet._number(scale.get("value"), 0))}', 9)
        if not columns:
            return
        computed = sheet.compute(block)
        grid = doc.add_table(rows=1 + len(rows), cols=len(columns))
        grid.style = 'Table Grid'
        for c, col in enumerate(columns):
            cell = grid.rows[0].cells[c]
            p = cell.paragraphs[0]
            self.grey(p, f'{sheet.col_name(c)} ', 7.5)
            p.add_run(col.get('name') or '').bold = True
            if col.get('unit'):
                self.grey(p, f' ({col["unit"]})', 8)
            if col.get('kind') == 'formula' and col.get('formula'):
                self.grey(cell.add_paragraph(), col['formula'], 7)
            if col.get('kind') == 'fill':
                self.grey(cell.add_paragraph(), 'filled in on the day', 7)
        for r, row in enumerate(rows):
            given = row.get('cells') or []
            for c, col in enumerate(columns):
                raw = given[c] if c < len(given) else None
                p = grid.rows[r + 1].cells[c].paragraphs[0]
                if col.get('kind') == 'check':
                    p.add_run('☑' if raw else '☐')
                    if isinstance(raw, dict) and raw.get('by'):
                        run = p.add_run(f' {raw["by"]} · {self.ctx.when(raw.get("at"))}')
                        run.font.size = self.Pt(7.5)
                        run.font.color.rgb = self.RGBColor(*GREEN)
                    continue
                shown = computed[r][c]
                text = raw if col.get('kind') == 'text' and not str(raw or '').startswith('=') else shown['display']
                run = p.add_run(str(text or ''))
                run.font.size = self.Pt(9)
                if shown.get('error'):
                    run.font.color.rgb = self.RGBColor(0xB9, 0x1C, 0x1C)
        doc.add_paragraph()

    def section(self, block: dict) -> None:
        src = block.get('source') or {}
        p = self.doc.add_paragraph()
        self.shade(p)
        p.add_run(block.get('title') or '').bold = True
        origin = (f'Protocol{f" · version {src["version"]}" if src.get("version") else ""}'
                  if src.get('kind') == 'protocol' else 'From the MoClo Library Tool' if src.get('kind') == 'moclo-reaction' else '')
        if origin:
            p.add_run('   ')
            if src.get('kind') == 'protocol' and src.get('id'):
                if src['id'] in self.ctx.included:
                    self.link(p, origin, anchor=f'n{src["id"][:38]}')
                else:
                    self.link(p, origin, url=self.ctx.note_link(src['id']))
            else:
                self.grey(p, origin, 8.5)
        for b in block.get('blocks') or []:
            self.block(b)
        self.grey(self.doc.add_paragraph(), '─' * 30, 7)

    def block(self, block: dict) -> None:
        kind = block.get('type')
        if kind == 'text':
            self.text(block)
        elif kind == 'steps':
            self.steps(block)
        elif kind == 'table':
            self.table(block)
        elif kind == 'section':
            self.section(block)

    # --- the document --------------------------------------------------------------------------

    def entry(self, e: Entry, breaks: bool) -> None:
        ctx, doc = self.ctx, self.doc
        self.mentioned = set()
        heading = doc.add_heading(e.title, level=min(e.depth + 1, 9))
        heading.paragraph_format.page_break_before = breaks
        self.mark(heading, f'n{e.id[:38]}')
        meta = [exporting.TYPE_NAMES.get(e.type, e.type)]
        if e.where:
            meta.append(e.where)
        if e.created_by:
            meta.append(f'created by {e.created_by} · {ctx.when(e.created_at)}')
        if e.updated_by and (e.updated_by != e.created_by or e.updated_at != e.created_at):
            meta.append(f'last saved by {e.updated_by} · {ctx.when(e.updated_at)}')
        if not e.created_by:
            meta.append('supplied with the Lab Notebook')
        meta.append(f'version {e.revision}')
        self.grey(doc.add_paragraph(), ' · '.join(meta))
        if e.repeat_of:
            p = doc.add_paragraph('This page is in more than one place; it is written out ')
            self.link(p, 'here', anchor=f'n{e.repeat_of[:38]}')
            p.add_run('.')
            return
        if e.labels:
            run = doc.add_paragraph().add_run('  '.join(e.labels))
            run.font.size = self.Pt(8.5)
            run.font.color.rgb = self.RGBColor(0x43, 0x38, 0xCA)
        start = len(doc.paragraphs)
        for b in e.content:
            self.block(b)
        files = [a for a in e.attachments if not a['deleted_at'] and a['id'] not in self.mentioned]
        if files:
            doc.add_paragraph().add_run('Files').bold = True
            for a in files:
                marker = len(doc.paragraphs)
                self.shown(a['id'])
                if len(doc.paragraphs) > marker:
                    self.mark(doc.paragraphs[marker], f'a{a["id"][:30]}')
        if len(doc.paragraphs) == start and not e.content and not e.has_children and e.type not in ('year', 'month'):
            run = doc.add_paragraph().add_run('Nothing written here.')
            run.italic = True

    def write(self, step: Callable) -> None:
        ctx, doc = self.ctx, self.doc
        doc.core_properties.title = ctx.title
        doc.core_properties.author = ctx.person.username
        doc.add_heading(ctx.title, level=0)
        self.grey(doc.add_paragraph(), ctx.subtitle, 11)
        pages = sum(1 for e in ctx.entries if not e.repeat_of)
        files = sum(len([a for a in e.attachments if not a['deleted_at']]) for e in ctx.entries)
        info = doc.add_paragraph()
        info.add_run(f'Exported by {ctx.person.username}{f" ({ctx.person.lab})" if ctx.person.lab else ""} on {ctx.stamp}. ')
        info.add_run(f'{pages} page{"s" if pages != 1 else ""}, {files} file{"s" if files != 1 else ""}. ')
        self.grey(info, 'From the Lab Notebook of the Gardiner Lab suite: a copy to edit. The notebook keeps '
                        'every version of each page; this is the current one.', 9)
        if len(ctx.entries) > 1:
            doc.add_paragraph().add_run('Contents').bold = True
            for e in ctx.entries:
                if e.repeat_of:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = self.Inches(0.2 * min(e.depth, 6))
                p.paragraph_format.space_after = self.Pt(0)
                self.link(p, e.title, anchor=f'n{e.id[:38]}')
        previous_had_content = True
        for i, e in enumerate(ctx.entries):
            self.entry(e, breaks=previous_had_content)
            previous_had_content = bool(e.content or e.attachments) or e.type not in ('year', 'month', 'folder')
            step(i + 1)
        footer = doc.sections[0].footer.paragraphs[0]
        self.grey(footer, f'{ctx.title} · exported {ctx.stamp} by {ctx.person.username}', 7.5)


def render(ctx: Context, step: Callable) -> Tuple[bytes, str, str]:
    writer = _Writer(ctx)
    writer.write(step)
    out = io.BytesIO()
    writer.doc.save(out)
    document = out.getvalue()
    if not writer.beside:
        return document, 'docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    step(len(ctx.entries), 'Putting the files beside it')
    name = exporting.safe_filename(f'{ctx.title} {ctx.stamp[:10]}', 'docx')
    bundle = io.BytesIO()
    with zipfile.ZipFile(bundle, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr(name, document)
        for path, data in writer.beside.values():
            z.writestr(path, data)
        z.writestr('README.txt', f'{ctx.title}\r\n\r\nOpen {name}. The files it links to are in the files folder: '
                                 'keep the two together (unzip the whole of it) so the links work.\r\n')
    return bundle.getvalue(), 'zip', 'application/zip'
