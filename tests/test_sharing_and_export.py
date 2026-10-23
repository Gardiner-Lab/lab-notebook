"""Protocols kept to their lab and shared with others; notebooks shared by lab leads; exports."""

import io
import zipfile

from app import exporting


def _me(client):
    return client.get('/api/me').get_json()


def _protocol(client, parent, title, content=None):
    r = client.post('/api/notes', json={'parent_id': parent, 'title': title, 'type': 'protocol', 'content': content})
    assert r.status_code == 201, r.get_json()
    return r.get_json()['note']


def _png(colour=(200, 30, 30)):
    from PIL import Image
    out = io.BytesIO()
    Image.new('RGB', (120, 80), colour).save(out, 'PNG')
    return out.getvalue()


def _pdf(pages=2):
    from pypdf import PdfWriter
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(300, 400)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def _attach(client, note_id, *files):
    data = {'file': [(io.BytesIO(body), name) for name, body in files]}
    r = client.post(f'/api/notes/{note_id}/attachments', data=data, content_type='multipart/form-data')
    assert r.status_code == 201, r.get_json()
    return r.get_json()['attachments']


def test_protocols_stay_in_their_lab_until_shared(as_user):
    ana, ben, lead, zoe, olead = (as_user(u) for u in ('ana', 'ben', 'lead', 'zoe', 'olead'))
    library = _me(ana)['library']
    folder = ana.post('/api/notes', json={'parent_id': library, 'title': 'Cloning', 'type': 'folder'}).get_json()['note']
    proto = _protocol(ana, folder['id'], 'Miniprep')
    secret = _protocol(ana, folder['id'], 'Our own trick')

    assert ben.get(f"/api/notes/{proto['id']}").status_code == 200           # the lab's own
    assert zoe.get(f"/api/notes/{proto['id']}").status_code == 403           # another lab's
    titles = lambda c: [p['title'] for p in c.get('/api/protocols').get_json()['protocols']]
    assert 'Miniprep' in titles(ben) and 'Miniprep' not in titles(zoe)
    # The ones supplied with the notebook are every lab's.
    assert 'MoClo Golden Gate assembly' in titles(zoe)
    assert [c['title'] for c in zoe.get(f'/api/notes/{library}/children').get_json()['children']] == \
        ['How to write a protocol (start here)', 'MoClo Golden Gate assembly']

    # The author shares it with the other lab; the other lab can use it, not change it.
    offer = ana.get(f"/api/notes/{proto['id']}/shares").get_json()
    assert offer['can_share'] and [lab['name'] for lab in offer['labs']] == ['Other Lab'] and offer['people'] == []
    shared = ana.post(f"/api/notes/{proto['id']}/shares", json={'target_type': 'lab', 'target': '2'})
    assert shared.status_code == 201 and [s['label'] for s in shared.get_json()['shares']] == ['Other Lab']
    assert zoe.get(f"/api/notes/{proto['id']}").get_json()['note']['can_edit'] is False
    assert zoe.put(f"/api/notes/{proto['id']}", json={'title': 'Mine now'}).status_code == 403
    assert zoe.get(f"/api/protocols/{proto['id']}/snapshot").status_code == 200
    listed = {p['title']: p for p in zoe.get('/api/protocols').get_json()['protocols']}
    assert listed['Miniprep']['shared'] is True and listed['Miniprep']['lab'] == 'Gardiner Lab'
    assert 'Our own trick' not in listed
    # The folder shows, with only what is shared in it; nothing can be added to it from there.
    assert [c['title'] for c in zoe.get(f"/api/notes/{folder['id']}/children").get_json()['children']] == ['Miniprep']
    assert zoe.post('/api/notes', json={'parent_id': folder['id'], 'title': 'x', 'type': 'protocol'}).status_code == 403
    assert zoe.get(f"/api/notes/{secret['id']}").status_code == 403

    # Only who may change it shares it: the author, their lead, an admin. Not the other lab.
    assert ben.post(f"/api/notes/{proto['id']}/shares", json={'target_type': 'lab', 'target': '2'}).status_code == 403
    assert olead.post(f"/api/notes/{proto['id']}/shares", json={'target_type': 'lab', 'target': '2'}).status_code == 403
    assert lead.get(f"/api/notes/{proto['id']}/shares").get_json()['can_share'] is True
    assert ana.post(f"/api/notes/{proto['id']}/shares", json={'target_type': 'lab', 'target': '1'}).status_code == 400
    assert ana.post(f"/api/notes/{proto['id']}/shares", json={'target_type': 'person', 'target': 'zoe'}).status_code == 400
    moclo = next(p for p in ana.get('/api/protocols').get_json()['protocols'] if p['title'].startswith('MoClo'))
    assert as_user('boss').post(f"/api/notes/{moclo['id']}/shares", json={'target_type': 'lab', 'target': '2'}).status_code == 400

    # Shared until retracted.
    share_id = shared.get_json()['shares'][0]['id']
    assert zoe.delete(f'/api/shares/{share_id}').status_code == 403
    assert lead.delete(f'/api/shares/{share_id}').status_code == 200
    assert zoe.get(f"/api/notes/{proto['id']}").status_code == 403


def test_protocols_written_before_labs_were_recorded_take_their_authors_lab(as_user, world):
    import sqlite3
    from app import notes
    ana, zoe = as_user('ana'), as_user('zoe')
    proto = _protocol(ana, _me(ana)['library'], 'Old protocol')
    conn = sqlite3.connect(world['paths']['notebook'])
    conn.execute('UPDATE notes SET lab = NULL WHERE id = ?', (proto['id'],))
    conn.commit()
    assert zoe.get(f"/api/notes/{proto['id']}").status_code == 403         # the author's lab, read now
    notes.assign_labs()
    assert conn.execute('SELECT lab FROM notes WHERE id = ?', (proto['id'],)).fetchone()[0] == '1'


def test_lab_leads_share_notebooks_with_labs_and_people(as_user):
    ana, ben, lead, zoe, olead = (as_user(u) for u in ('ana', 'ben', 'lead', 'zoe', 'olead'))
    day = ana.post('/api/days', json={'date': '2026-10-16'}).get_json()['note']
    notebook = _me(ana)['notebook']
    assert _me(ana)['can_share_notebook'] is False and _me(lead)['can_share_notebook'] is True

    # A member does not share notebooks, not even their own; the lab's lead does.
    assert ana.post(f'/api/notes/{notebook}/shares', json={'target_type': 'person', 'target': 'zoe'}).status_code == 403
    assert olead.post(f'/api/notes/{notebook}/shares', json={'target_type': 'person', 'target': 'zoe'}).status_code == 403
    offer = lead.get(f'/api/notes/{notebook}/shares').get_json()
    assert 'zoe' in offer['people'] and 'ana' not in offer['people'] and [lab['id'] for lab in offer['labs']] == ['2']
    r = lead.post(f'/api/notes/{notebook}/shares', json={'target_type': 'person', 'target': 'zoe'})
    assert r.status_code == 201

    # Zoe reads it (and its files), as a lab lead would; only ana changes it.
    seen = zoe.get(f"/api/notes/{day['id']}")
    assert seen.status_code == 200 and seen.get_json()['note']['can_edit'] is False
    assert zoe.put(f"/api/notes/{day['id']}", json={'content': []}).status_code == 403
    assert [(n['owner'], n['shared']) for n in _me(zoe)['others']] == [('ana', True)]
    assert olead.get(f"/api/notes/{day['id']}").status_code == 403
    assert ben.get(f"/api/notes/{day['id']}").status_code == 403
    assert [n['title'] for n in zoe.get('/api/search?q=2026-10-16').get_json()['notes']] == ['2026-10-16 Friday']

    # With a lab: everyone in it.
    lab_share = lead.post(f'/api/notes/{notebook}/shares', json={'target_type': 'lab', 'target': '2'}).get_json()
    assert olead.get(f"/api/notes/{day['id']}").status_code == 200
    assert {s['label'] for s in lab_share['shares']} == {'zoe', 'Other Lab'}
    # A day is not shared alone: notebooks are shared whole.
    assert lead.post(f"/api/notes/{day['id']}/shares", json={'target_type': 'lab', 'target': '2'}).status_code == 403

    for s in lab_share['shares']:
        assert lead.delete(f"/api/shares/{s['id']}").status_code == 200
    assert zoe.get(f"/api/notes/{day['id']}").status_code == 403
    assert _me(zoe)['others'] == []


def _export(client, note_id, fmt='pdf', **extra):
    r = client.post(f'/api/notes/{note_id}/export', json={'format': fmt, 'timezone': 'Europe/Amsterdam',
                                                          'origin': 'http://lab.example:8080', **extra})
    assert r.status_code == 202, r.get_json()
    job = r.get_json()['job']
    exporting.wait(job['id'])
    status = client.get(f"/api/exports/{job['id']}").get_json()['job']
    assert status['status'] == 'done', status
    got = client.get(f"/api/exports/{job['id']}/file")
    assert got.status_code == 200
    return job, status, got


def _notebook_with_everything(ana):
    notebook = _me(ana)['notebook']
    day = ana.post('/api/days', json={'date': '2026-10-16'}).get_json()['note']
    gel, scan, plate = _attach(ana, day['id'], ('gel.png', _png()), ('scan.pdf', _pdf(2)),
                               ('plate.csv', b'well,od\nA1,0.51\nA2,0.48\n'))
    word = _attach(ana, day['id'], ('notes.docx', b'PK\x03\x04 not really a document'))[0]
    content = [
        {'type': 'text', 'html': f'<h3>Miniprep</h3><p>Grew <b>pGFP</b> in <a data-plasmid="p1">pGFP-35S</a>.</p>'
                                 f'<p><img data-attachment="{gel["id"]}" alt="gel"></p>'
                                 f'<p>Scan: <a data-attachment="{scan["id"]}">📕 scan.pdf</a></p>'
                                 '<ul><li>one</li><li>two<ol><li>nested</li></ol></li></ul>'},
        {'type': 'steps', 'title': 'Steps', 'items': [{'text': 'Spin down', 'done': True}, {'text': 'Elute'}]},
        {'type': 'table', 'title': 'Mix', 'scale': {'label': 'Reactions', 'value': 3, 'base': 1},
         'columns': [{'name': 'Component', 'kind': 'text'}, {'name': 'Per reaction', 'kind': 'number', 'unit': 'µL'},
                     {'name': 'Total', 'kind': 'formula', 'unit': 'µL', 'formula': '=B*N'}, {'name': 'Added', 'kind': 'check'}],
         'rows': [{'cells': ['Buffer', '2.5', '', True]}, {'cells': ['Water', '=20-B1', '', None]}]},
    ]
    saved = ana.put(f"/api/notes/{day['id']}", json={'content': content}).get_json()['note']
    child = ana.post('/api/notes', json={'parent_id': day['id'], 'title': 'Gel analysis',
                                         'content': [{'type': 'text', 'html': '<p>Bands at 1 kb.</p>'}]}).get_json()['note']
    return notebook, saved, child, {'gel': gel, 'scan': scan, 'plate': plate, 'word': word}


def test_a_pdf_export_has_every_page_under_it_with_its_pictures_and_files(as_user):
    from pypdf import PdfReader
    ana = as_user('ana')
    notebook, day, child, files = _notebook_with_everything(ana)

    job, status, got = _export(ana, day['id'])
    assert got.mimetype == 'application/pdf' and status['filename'].endswith('.pdf')
    assert got.headers['Content-Disposition'].startswith('attachment')
    pdf = PdfReader(io.BytesIO(got.data))
    text = '\n'.join(p.extract_text() for p in pdf.pages)
    for words in ('2026-10-16 Friday', 'Gel analysis', 'Bands at 1 kb', 'Spin down', 'Reactions: 3',
                  '7.5', '52.5', 'scan.pdf · page 2 of 2', 'A1', '0.51', 'Contents'):
        assert words in text, words
    # Times in the exporter's zone; the tick recorded by the server.
    assert '✓ ana ·' in text
    # Its two attached pages each have a page of their own, drawn onto it.
    captions = [i for i, p in enumerate(pdf.pages) if 'scan.pdf · page' in p.extract_text()]
    assert len(captions) == 2 and captions[1] == captions[0] + 1
    # Other files are embedded, to be opened from the PDF.
    annotations = [a.get_object() for p in pdf.pages for a in (p.get('/Annots') or [])]
    embedded = {a['/FS'].get_object().get('/UF') or a['/FS'].get_object().get('/F')
                for a in annotations if a.get('/Subtype') == '/FileAttachment'}
    assert {'scan.pdf', 'plate.csv', 'notes.docx'} <= embedded
    # Links go to the suite's tools, from where it was exported.
    links = {a['/A'].get('/URI') for a in annotations if a.get('/Subtype') == '/Link' and '/URI' in a.get('/A', {})}
    # The photo is shown once, where it is in the text, not again under the files.
    assert text.count('gel.png') == 0 and 'Files' in text
    assert 'http://lab.example:8080/editor/editor?open=p1' in links

    # Only the page itself, without what is under it.
    _, _, alone = _export(ana, day['id'], children=False)
    assert 'Gel analysis' not in '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(alone.data)).pages)

    # The whole notebook, from its top.
    _, _, whole = _export(ana, notebook)
    whole_text = '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(whole.data)).pages)
    assert "ana's notebook" in whole_text and '2026-10 October' in whole_text and 'Bands at 1 kb' in whole_text


def test_a_word_export_is_editable_with_its_files_beside_it(as_user):
    import docx
    ana = as_user('ana')
    _, day, child, files = _notebook_with_everything(ana)
    _, status, got = _export(ana, day['id'], 'docx')
    assert got.mimetype == 'application/zip' and status['filename'].endswith('.zip')
    bundle = zipfile.ZipFile(io.BytesIO(got.data))
    names = bundle.namelist()
    assert {'files/scan.pdf', 'files/plate.csv', 'files/notes.docx', 'README.txt'} <= set(names)
    document = docx.Document(io.BytesIO(bundle.read(next(n for n in names if n.endswith('.docx')))))
    text = '\n'.join(p.text for p in document.paragraphs)
    for words in ('2026-10-16 Friday', 'Gel analysis', 'Bands at 1 kb', '☑ 1. Spin down', '• one', '1. nested',
                  'Reactions: 3'):
        assert words in text, words
    cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
    assert '7.5' in cells and '52.5' in cells and 'Buffer' in cells
    headings = [p.text for p in document.paragraphs if p.style.name.startswith('Heading')]
    assert headings[:2] == ['2026-10-16 Friday', 'Gel analysis']
    # The photo, and the attached PDF's two pages, are pictures in the document.
    assert len(document.inline_shapes) == 3

    # Without other files it is the document alone.
    page = ana.post('/api/notes', json={'parent_id': _me(ana)['notebook'], 'title': 'Plain',
                                        'content': [{'type': 'text', 'html': '<p>Just text</p>'}]}).get_json()['note']
    _, plain_status, plain = _export(ana, page['id'], 'docx')
    assert plain_status['filename'].endswith('.docx') and plain.data[:2] == b'PK'
    assert 'Just text' in '\n'.join(p.text for p in docx.Document(io.BytesIO(plain.data)).paragraphs)


def test_who_may_export_what(as_user):
    from pypdf import PdfReader
    ana, ben, lead, zoe = (as_user(u) for u in ('ana', 'ben', 'lead', 'zoe'))
    day = ana.post('/api/days', json={'date': '2026-10-16'}).get_json()['note']
    assert zoe.post(f"/api/notes/{day['id']}/export", json={'format': 'pdf'}).status_code == 403
    assert ana.post(f"/api/notes/{day['id']}/export", json={'format': 'odt'}).status_code == 400
    # A lead exports their lab's notebooks; an export is its maker's only.
    job, _, _ = _export(lead, day['id'])
    assert ana.get(f"/api/exports/{job['id']}").status_code == 404
    assert ana.get(f"/api/exports/{job['id']}/file").status_code == 404

    # A protocol, and the library as each lab sees it.
    library = _me(ana)['library']
    _protocol(ana, library, 'Western blot', [{'type': 'steps', 'items': [{'text': 'Block'}]}])
    _, _, theirs = _export(zoe, library)
    text = '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(theirs.data)).pages)
    assert 'MoClo Golden Gate assembly' in text and 'Western blot' not in text
    _, _, ours = _export(ben, library)
    assert 'Western blot' in '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(ours.data)).pages)


def test_a_gel_photo_with_its_wells_marked(as_user):
    import docx
    from pypdf import PdfReader
    ana = as_user('ana')
    page = ana.post('/api/notes', json={'parent_id': _me(ana)['notebook'], 'title': 'Colony PCR'}).get_json()['note']
    photo = _attach(ana, page['id'], ('gel.png', _png((20, 20, 20))))[0]
    # Clicked out of order; one click off the photo, one with a made-up attachment.
    gel = {'type': 'gel', 'title': '1% agarose', 'attachment': photo['id'], 'lanes': [
        {'x': 0.62, 'y': 0.1, 'label': 'pXK17 clone 2', 'note': '5 µL'},
        {'x': 0.18, 'y': 0.12, 'label': 'ladder'},
        {'x': 1.7, 'y': 1.4, 'label': 'clone 3'}]}
    saved = ana.put(f"/api/notes/{page['id']}", json={'content': [gel, {'type': 'gel', 'attachment': '../etc/passwd', 'lanes': []}]})
    assert saved.status_code == 200, saved.get_json()
    first, second = saved.get_json()['note']['content']
    assert [lane['label'] for lane in first['lanes']] == ['ladder', 'pXK17 clone 2', 'clone 3'], 'top row left to right, then the next'
    assert (first['lanes'][2]['x'], first['lanes'][2]['y']) == (1.0, 1.0)
    assert first['attachment'] == photo['id'] and second['attachment'] == ''
    assert [n['title'] for n in ana.get('/api/search?q=pXK17').get_json()['notes']] == ['Colony PCR']

    # A copy has its own copy of the photo.
    copy = ana.post(f"/api/notes/{page['id']}/copy").get_json()['note']
    assert copy['content'][0]['attachment'] not in ('', photo['id'])
    assert copy['content'][0]['attachment'] == copy['attachments'][0]['id']

    # Exports show the photo once, with its numbers, and the lanes beside it.
    _, _, pdf = _export(ana, page['id'], children=False)
    text = '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.data)).pages)
    assert 'ladder' in text and 'clone 2' in text and '5 µL' in text and 'Lane' in text
    assert 'Files' not in text, 'the photo is the gel\'s, not listed again'
    _, _, word = _export(ana, page['id'], 'docx', children=False)
    document = docx.Document(io.BytesIO(word.data))
    lanes = [c.text for outer in document.tables for cell in outer._cells for inner in cell.tables
             for row in inner.rows for c in row.cells]
    assert lanes[:6] == ['Lane', 'Sample', 'Notes', '1', 'ladder', ''] and 'pXK17 clone 2' in lanes
    assert len(document.inline_shapes) == 1


def test_numbers_are_drawn_on_the_gel_for_exports():
    from PIL import Image
    from app import exporting
    picture = exporting.gel_picture(_png((0, 0, 0)), [{'x': 0.5, 'y': 0.5}, {'x': 0.1, 'y': 0.5}])
    image = Image.open(io.BytesIO(picture)).convert('RGB')
    assert image.size == (120, 80)
    assert image.getpixel((60, 40)) != (0, 0, 0) and image.getpixel((12, 40)) != (0, 0, 0)
    assert exporting.gel_picture(b'not a picture', []) is None


def test_a_gel_with_two_rows_of_wells_is_numbered_row_by_row(as_user):
    from app.content import gel_rows
    ana = as_user('ana')
    page = ana.post('/api/notes', json={'parent_id': _me(ana)['notebook'], 'title': 'Two combs'}).get_json()['note']
    photo = _attach(ana, page['id'], ('gel.png', _png((20, 20, 20))))[0]
    # Clicked in no order; each row a little crooked, as wells are on a photo.
    clicks = [('B2', 0.40, 0.56), ('A3', 0.70, 0.11), ('A1', 0.10, 0.10), ('B1', 0.12, 0.53),
              ('A2', 0.40, 0.13), ('B3', 0.71, 0.55)]
    saved = ana.put(f"/api/notes/{page['id']}", json={'content': [{'type': 'gel', 'attachment': photo['id'], 'lanes': [
        {'x': x, 'y': y, 'label': name} for name, x, y in clicks]}]}).get_json()['note']
    assert [lane['label'] for lane in saved['content'][0]['lanes']] == ['A1', 'A2', 'A3', 'B1', 'B2', 'B3']
    assert [[lane['label'] for lane in row] for row in gel_rows(saved['content'][0]['lanes'])] == \
        [['A1', 'A2', 'A3'], ['B1', 'B2', 'B3']]
    # One row of wells is still left to right, however crooked within the gap.
    assert [lane['x'] for lane in gel_rows([{'x': 0.5, 'y': 0.12}, {'x': 0.2, 'y': 0.08}, {'x': 0.8, 'y': 0.1}])[0]] == [0.2, 0.5, 0.8]


def test_plants_on_a_photo_are_numbered_in_the_order_clicked(as_user):
    import docx
    from pypdf import PdfReader
    ana = as_user('ana')
    page = ana.post('/api/notes', json={'parent_id': _me(ana)['notebook'], 'title': 'T1 selection'}).get_json()['note']
    photo = _attach(ana, page['id'], ('tray.png', _png((40, 90, 40))))[0]
    clicked = [('Col-0 #1', 0.8, 0.7), ('Col-0 #2', 0.1, 0.2), ('Col-0 #3', 0.5, 0.5)]
    block = {'type': 'gel', 'order': 'clicks', 'attachment': photo['id'],
             'lanes': [{'x': x, 'y': y, 'label': name, 'note': 'resistant' if n == 1 else ''} for n, (name, x, y) in enumerate(clicked)]}
    saved = ana.put(f"/api/notes/{page['id']}", json={'content': [block]}).get_json()['note']['content'][0]
    assert saved['order'] == 'clicks'
    assert [lane['label'] for lane in saved['lanes']] == ['Col-0 #1', 'Col-0 #2', 'Col-0 #3'], 'as clicked, not by place'
    # Anything but 'clicks' is a gel's numbering, by rows.
    gel = ana.put(f"/api/notes/{page['id']}", json={'content': [{**block, 'order': 'nonsense'}]}).get_json()['note']['content'][0]
    assert gel['order'] == 'rows' and [lane['label'] for lane in gel['lanes']] == ['Col-0 #2', 'Col-0 #3', 'Col-0 #1']

    ana.put(f"/api/notes/{page['id']}", json={'content': [block]})
    _, _, pdf = _export(ana, page['id'], children=False)
    text = '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.data)).pages)
    assert 'No.' in text and 'Label' in text and 'Lane' not in text
    assert text.index('Col-0 #1') < text.index('Col-0 #2') < text.index('Col-0 #3')
    _, _, word = _export(ana, page['id'], 'docx', children=False)
    document = docx.Document(io.BytesIO(word.data))
    cells = [c.text for outer in document.tables for cell in outer._cells for inner in cell.tables
             for row in inner.rows for c in row.cells]
    assert cells[:9] == ['No.', 'Label', 'Notes', '1', 'Col-0 #1', '', '2', 'Col-0 #2', 'resistant']


def test_someone_in_two_labs_sees_and_leads_in_each_as_their_role_there(as_user, world):
    import sqlite3
    # Zoe joins Gardiner Lab as well, and leads it; she stays a member of Other Lab.
    conn = sqlite3.connect(world['paths']['suite'])
    conn.executescript("""
        CREATE TABLE lab_members (username TEXT, lab_id INTEGER, role TEXT, lab_admin INTEGER DEFAULT 0,
                                  joined_at TEXT, PRIMARY KEY (username, lab_id));
        INSERT INTO lab_members (username, lab_id, role, joined_at) VALUES
            ('ana', 1, 'member', '1'), ('ben', 1, 'member', '2'), ('lead', 1, 'lab_lead', '3'),
            ('zoe', 2, 'member', '4'), ('olead', 2, 'lab_lead', '5'), ('zoe', 1, 'lab_lead', '6');
    """)
    conn.commit()
    ana, zoe, olead = as_user('ana'), as_user('zoe'), as_user('olead')
    proto = _protocol(ana, _me(ana)['library'], 'Gardiner miniprep')
    # Gardiner Lab's protocol: Zoe sees it, and, leading Gardiner Lab, may change it.
    seen = zoe.get(f"/api/notes/{proto['id']}")
    assert seen.status_code == 200 and seen.get_json()['note']['can_edit'] is True
    assert olead.get(f"/api/notes/{proto['id']}").status_code == 403
    # Leading Gardiner Lab, she reads its members' notebooks; Other Lab's lead does not.
    day = ana.post('/api/days', json={'date': '2026-10-17'}).get_json()['note']
    assert zoe.get(f"/api/notes/{day['id']}").status_code == 200
    assert olead.get(f"/api/notes/{day['id']}").status_code == 403
    # Shares with either of her labs reach her.
    assert [m['lab'] for m in zoe.get('/api/me').get_json()['person']['labs']] == ['Other Lab', 'Gardiner Lab']


def test_a_virtual_digest_goes_into_the_day(as_user):
    """The Plasmid Editor puts a virtual digest - its gel and fragments - into a day's page."""
    import json
    ana, zoe = as_user('ana'), as_user('zoe')
    section = {
        'type': 'section', 'title': 'Virtual digest: pGFP-35S (EcoRI+BamHI)',
        'source': {'kind': 'virtual-digest', 'id': '', 'name': 'pGFP-35S'},
        'blocks': [
            {'type': 'text', 'html': '<p>Digest of <a data-plasmid="p1">pGFP-35S</a>, beside a 1 kb ladder.</p>'},
            {'type': 'gel', 'title': 'Virtual gel', 'attachment': 'picture', 'order': 'rows',
             'lanes': [{'x': 0.1, 'y': 0.05, 'label': 'Ladder: 1 kb ladder'},
                       {'x': 0.3, 'y': 0.05, 'label': 'pGFP-35S · EcoRI+BamHI', 'note': '1,200 + 400 bp'}]},
            {'type': 'table', 'title': 'Fragments',
             'columns': [{'name': 'Lane', 'kind': 'text'}, {'name': 'Length', 'kind': 'number', 'unit': 'bp'},
                         {'name': 'Seen on the gel', 'kind': 'check'}],
             'rows': [{'cells': ['2', '1200', None]}, {'cells': ['2', '400', None]}]}]}
    form = lambda **extra: {'section': json.dumps(section), **extra}

    r = ana.post('/api/days/import', data={**form(date='2026-03-04'), 'picture': (io.BytesIO(_png()), 'gel.png')},
                 content_type='multipart/form-data')
    assert r.status_code == 201, r.get_json()
    day = r.get_json()['note']
    assert day['type'] == 'day' and day['title'].startswith('2026-03-04')
    # It takes the place of the empty paragraph a new day starts with.
    assert [b['type'] for b in day['content']] == ['section']
    got = day['content'][0]
    assert got['source']['kind'] == 'virtual-digest'
    gel = got['blocks'][1]
    assert gel['attachment'] in [a['id'] for a in day['attachments']]
    assert [lane['label'] for lane in gel['lanes']] == ['Ladder: 1 kb ladder', 'pGFP-35S · EcoRI+BamHI']
    assert ana.get(f"/api/attachments/{gel['attachment']}").status_code == 200
    # Linked from the plasmid, found by its words, kept in the history as an import.
    assert ('plasmid', 'p1') in [(t, i) for t, i in __import__('app').content.links(day['content'])]
    assert day['id'] in [n['id'] for n in ana.get('/api/search?q=EcoRI').get_json()['notes']]
    history = ana.get(f"/api/notes/{day['id']}/revisions").get_json()['revisions']
    assert history[0]['kind'] == 'imported' and 'from the Plasmid Editor' in history[0]['note']

    # A second one goes after it, on the same day; what was written there stays.
    again = ana.post('/api/days/import', data=form(date='2026-03-04'), content_type='multipart/form-data').get_json()['note']
    assert again['id'] == day['id'] and [b['type'] for b in again['content']] == ['section', 'section']

    # Into one's own notebook only; and only a virtual digest.
    assert zoe.get(f"/api/notes/{day['id']}").status_code == 403
    bad = dict(section, source={'kind': 'protocol'})
    assert ana.post('/api/days/import', data={'section': json.dumps(bad)},
                    content_type='multipart/form-data').status_code == 400
    assert ana.post('/api/days/import', data={'section': 'nonsense'}, content_type='multipart/form-data').status_code == 400
